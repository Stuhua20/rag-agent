import asyncio
import tempfile
import unittest
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi import UploadFile

from app.core.config import settings
from app.services.document_pipeline import DocumentPipeline
from app.services.json_store_service import json_store
from app.services.pdf_parser_service import PdfParserService
from app.services.table_chunker_service import TableChunkerService
from app.services.word_parser_service import WordParserService


class DocumentParserTests(unittest.TestCase):
    def test_pdf_two_column_text_is_ordered_by_column(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "two-column.pdf"
            pdf = fitz.open()
            page = pdf.new_page(width=600, height=800)
            page.insert_text((50, 50), "Article title")
            for index, y_position in enumerate((110, 145, 180, 215), start=1):
                page.insert_text((50, y_position), f"Left paragraph {index}")
                page.insert_text((330, y_position), f"Right paragraph {index}")
            pdf.save(path)
            pdf.close()

            parsed = PdfParserService().parse(path, "doc_pdf", "two-column.pdf")

        texts = [row["values"]["文本"] for row in parsed["sheets"][0]["rows"]]
        self.assertEqual(texts[0], "Article title")
        self.assertEqual(texts[1:5], ["Left paragraph 1", "Left paragraph 2", "Left paragraph 3", "Left paragraph 4"])
        self.assertEqual(texts[5:], ["Right paragraph 1", "Right paragraph 2", "Right paragraph 3", "Right paragraph 4"])

    def test_pdf_without_text_reports_unsupported_content(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "empty.pdf"
            pdf = fitz.open()
            pdf.new_page(width=600, height=800)
            pdf.save(path)
            pdf.close()

            with self.assertRaisesRegex(ValueError, "Scanned or image-only PDFs"):
                PdfParserService().parse(path, "doc_empty", "empty.pdf")

    def test_pdf_upload_flows_through_indexing_pipeline(self) -> None:
        original_data_dir = json_store.data_dir
        original_storage_dir = settings.storage_dir
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            json_store.data_dir = root / "data"
            settings.storage_dir = root / "storage"
            json_store.initialize()

            pdf_path = root / "pipeline.pdf"
            pdf = fitz.open()
            page = pdf.new_page(width=600, height=800)
            page.insert_text((50, 100), "Indexed PDF content")
            pdf.save(pdf_path)
            pdf.close()

            try:
                with pdf_path.open("rb") as file:
                    result = asyncio.run(
                        DocumentPipeline().upload_and_parse(
                            UploadFile(file=file, filename=pdf_path.name),
                            pdf_path.name,
                            1,
                        )
                    )
                self.assertEqual(result["status"], "indexed")
                documents = json_store.list_items("documents.json")
                self.assertEqual(documents[0]["parse_warnings"], [])
                chunks = json_store.list_items("chunks.json")
                self.assertEqual(chunks[0]["chunk_type"], "document_blocks")
                self.assertIn("Indexed PDF content", chunks[0]["content"])
            finally:
                json_store.data_dir = original_data_dir
                settings.storage_dir = original_storage_dir

    def test_word_parser_preserves_paragraphs_and_detects_columns(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "two-column.docx"
            document = Document()
            document.add_paragraph("第一段正文")
            columns = document.sections[0]._sectPr.find(qn("w:cols"))
            if columns is None:
                columns = OxmlElement("w:cols")
                document.sections[0]._sectPr.append(columns)
            columns.set(qn("w:num"), "2")
            document.add_paragraph("第二段正文")
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "字段"
            table.rows[0].cells[1].text = "值"
            document.save(path)

            parsed = WordParserService().parse(path, "doc_word", "two-column.docx")

        self.assertEqual(parsed["layout"]["column_count"], 2)
        self.assertEqual(parsed["warnings"], [])
        texts = [row["values"]["文本"] for row in parsed["sheets"][0]["rows"]]
        self.assertEqual(texts[:2], ["第一段正文", "第二段正文"])
        self.assertIn("表格1 第1行", texts[2])

    def test_document_chunks_do_not_use_table_row_prefix(self) -> None:
        workbook = {
            "content_type": "document",
            "sheets": [
                {
                    "sheet_name": "第1页",
                    "headers": ["文本"],
                    "content_type": "document",
                    "rows": [
                        {"row_number": 1, "values": {"文本": "第一段"}},
                        {"row_number": 2, "values": {"文本": "第二段"}},
                    ],
                }
            ],
        }
        chunks = TableChunkerService().chunk_workbook(
            workbook,
            {"id": "doc", "file_name": "sample.pdf", "permission_level": 1},
            {"第1页": "sheet"},
            chunk_size=1000,
            overlap=0,
        )

        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0]["chunk_type"], "document_blocks")
        self.assertIn("第一段", chunks[0]["content"])
        self.assertNotIn("行 1：", chunks[0]["content"])

    def test_table_chunks_keep_existing_row_format(self) -> None:
        workbook = {
            "content_type": "table",
            "sheets": [
                {
                    "sheet_name": "销售",
                    "headers": ["客户", "金额"],
                    "content_type": "table",
                    "rows": [{"row_number": 2, "values": {"客户": "甲", "金额": 100}}],
                }
            ],
        }
        chunks = TableChunkerService().chunk_workbook(
            workbook,
            {"id": "doc", "file_name": "sample.csv", "permission_level": 1},
            {"销售": "sheet"},
            chunk_size=1000,
            overlap=0,
        )

        self.assertEqual(chunks[0]["chunk_type"], "table_rows")
        self.assertIn("行 2：客户=甲；金额=100", chunks[0]["content"])


if __name__ == "__main__":
    unittest.main()
