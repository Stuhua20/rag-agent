from pathlib import Path

from openpyxl import load_workbook


class ExcelParserService:
    def parse(self, file_path: str | Path, document_id: str, file_name: str) -> dict:
        path = Path(file_path)
        if path.suffix.lower() == ".xls":
            raise ValueError(".xls support is not enabled in this MVP. Please save the file as .xlsx or CSV.")

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheets = []
        for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
            rows = list(worksheet.iter_rows(values_only=True))
            header_position = self._find_header_row(rows)
            if header_position is None:
                continue
            raw_headers = list(rows[header_position])
            active_columns = [
                index for index, value in enumerate(raw_headers)
                if value is not None and str(value).strip() != ""
            ]
            headers = [str(raw_headers[index]).strip() for index in active_columns]
            parsed_rows = []
            for row_offset, raw_row in enumerate(rows[header_position + 1 :], start=header_position + 2):
                values = {}
                has_value = False
                for header, column_index in zip(headers, active_columns):
                    value = raw_row[column_index] if column_index < len(raw_row) else None
                    if value is not None and str(value).strip() != "":
                        has_value = True
                    values[header] = value
                if has_value:
                    parsed_rows.append({"row_number": row_offset, "values": values})

            sheets.append(
                {
                    "sheet_name": worksheet.title,
                    "sheet_index": sheet_index,
                    "header_row": header_position + 1,
                    "headers": headers,
                    "rows": parsed_rows,
                }
            )

        if not sheets:
            raise ValueError("No non-empty sheets with headers were found.")
        return {"document_id": document_id, "file_name": file_name, "sheets": sheets}

    def _find_header_row(self, rows: list[tuple]) -> int | None:
        for index, row in enumerate(rows):
            non_empty = [value for value in row if value is not None and str(value).strip() != ""]
            if non_empty:
                return index
        return None

