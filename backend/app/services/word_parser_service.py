from collections.abc import Iterator
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph


class WordParserService:
    """Extract body paragraphs and table rows from .docx documents."""

    def parse(self, file_path: str | Path, document_id: str, file_name: str) -> dict:
        path = Path(file_path)
        try:
            document = Document(str(path))
        except Exception as error:
            raise ValueError(f"Unable to read Word document: {error}") from error

        lines = self._extract_lines(document)
        if not lines:
            raise ValueError("Word document contains no extractable text.")

        return {
            "document_id": document_id,
            "file_name": file_name,
            "content_type": "document",
            "layout": {"column_count": self._column_count(document)},
            "warnings": self._layout_warnings(document),
            "sheets": [
                {
                    "sheet_name": "正文",
                    "sheet_index": 1,
                    "header_row": 1,
                    "headers": ["文本"],
                    "content_type": "document",
                    "rows": [
                        {"row_number": line_number, "values": {"文本": line}}
                        for line_number, line in enumerate(lines, start=1)
                    ],
                }
            ],
        }

    def _extract_lines(self, document: DocumentObject) -> list[str]:
        lines = []
        table_index = 0
        for block in self._iter_blocks(document):
            if isinstance(block, Paragraph):
                text = self._clean_text(block.text)
                if text:
                    lines.append(text)
                continue

            table_index += 1
            for row_index, row in enumerate(block.rows, start=1):
                cells = [self._clean_text(cell.text) for cell in row.cells]
                cells = [f"列{index + 1}={value}" for index, value in enumerate(cells) if value]
                if cells:
                    lines.append(f"表格{table_index} 第{row_index}行：" + "；".join(cells))
        return lines

    def _column_count(self, document: DocumentObject) -> int:
        counts = [1]
        for section in document.sections:
            columns = section._sectPr.find(qn("w:cols"))
            if columns is None:
                continue
            raw_count = columns.get(qn("w:num"))
            if raw_count:
                try:
                    counts.append(max(int(raw_count), 1))
                    continue
                except ValueError:
                    pass
            counts.append(max(len(columns.findall(qn("w:col"))), 1))
        return max(counts)

    def _layout_warnings(self, document: DocumentObject) -> list[str]:
        warnings = []
        if any(document.element.body.iter(qn("w:txbxContent"))):
            warnings.append("检测到文本框内容；当前版本只保证正文段落和表格，文本框文字可能未完整提取。")
        if self._has_header_footer_text(document):
            warnings.append("检测到页眉或页脚文字；当前版本只索引正文，页眉页脚不会单独加入正文。")
        return warnings

    def _has_header_footer_text(self, document: DocumentObject) -> bool:
        for section in document.sections:
            parts = (
                section.header,
                section.first_page_header,
                section.even_page_header,
                section.footer,
                section.first_page_footer,
                section.even_page_footer,
            )
            for part in parts:
                if any(self._clean_text(paragraph.text) for paragraph in part.paragraphs):
                    return True
                if any(
                    self._clean_text(cell.text)
                    for table in part.tables
                    for row in table.rows
                    for cell in row.cells
                ):
                    return True
        return False

    def _iter_blocks(self, document: DocumentObject) -> Iterator[Paragraph | Table]:
        for child in document.element.body.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, document)
            elif isinstance(child, CT_Tbl):
                yield Table(child, document)

    def _clean_text(self, value: str) -> str:
        return " ".join(value.replace("\r", "\n").replace("\n", " ").split())
